import os
import io
import logging
from fastapi import FastAPI, UploadFile, File, HTTPException, Form
from fastapi.middleware.cors import CORSMiddleware
from dotenv import load_dotenv
import pymupdf 
from pptx import Presentation 
from langchain_text_splitters import RecursiveCharacterTextSplitter
from supabase import create_client, Client
from google import genai
from google.genai import types
from openai import OpenAI
from pydantic import BaseModel
from typing import Optional
import docx

class ChatRequest(BaseModel):
    user_email: str
    message: str
    chat_id: Optional[str] = None 

class DeleteRequest(BaseModel):
    user_email: str

load_dotenv()
app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=[
        "http://localhost:5173", 
        "http://127.0.0.1:5173",
        "https://goofy-five.vercel.app" 
    ],
    allow_credentials=True,
    allow_methods=["*"],
    allow_headers=["*"],
)

# Initialize Supabase, Gemini (Primary + Embeddings), and OpenRouter (Fallback)
supabase: Client = create_client(
    os.getenv("SUPABASE_URL"), 
    os.getenv("SUPABASE_SERVICE_KEY")
)
gemini_client = genai.Client()
openrouter_client = OpenAI(
    api_key=os.getenv("OPENROUTER_API_KEY"),
    base_url="https://openrouter.ai/api/v1"
)

def extract_text_from_pdf(file_bytes: bytes) -> str:
    doc = pymupdf.open(stream=file_bytes, filetype="pdf")
    return "\n".join([page.get_text() for page in doc])

def extract_text_from_pptx(file_bytes: bytes) -> str:
    prs = Presentation(io.BytesIO(file_bytes))
    text = ""
    for slide in prs.slides:
        for shape in slide.shapes:
            if hasattr(shape, "text"):
                text += shape.text + "\n"
    return text

def extract_text_from_docx(file_bytes: bytes) -> str:
    doc = docx.Document(io.BytesIO(file_bytes))
    return "\n".join([paragraph.text for paragraph in doc.paragraphs])

@app.post("/upload")
async def upload_document(file: UploadFile = File(...), user_email: str = Form(...)):
    if not file.filename.lower().endswith(('.pdf', '.pptx', '.doc', '.docx')):
        raise HTTPException(status_code=400, detail="Only PDF, PPTX, and DOCX supported.")
    
    file_bytes = await file.read()
    
    if file.filename.lower().endswith('.pdf'):
        raw_text = extract_text_from_pdf(file_bytes)
    elif file.filename.lower().endswith('.pptx'):
        raw_text = extract_text_from_pptx(file_bytes)
    elif file.filename.lower().endswith(('.doc', '.docx')):
        raw_text = extract_text_from_docx(file_bytes)
    else:
        raise HTTPException(status_code=400, detail="Unsupported file format")
        
    text_splitter = RecursiveCharacterTextSplitter(
        chunk_size=1000,
        chunk_overlap=200
    )
    chunks = text_splitter.split_text(raw_text)
    
    records = []
    for chunk in chunks:
        result = gemini_client.models.embed_content(
            model="gemini-embedding-001",
            contents=chunk,
            config=types.EmbedContentConfig(output_dimensionality=768)
        )
        records.append({
            "filename": file.filename,
            "content": chunk,
            "embedding": result.embeddings[0].values,
            "user_email": user_email
        })
        
    if records:
         supabase.table("documents").insert(records).execute()
         
    return {"message": f"Successfully processed {len(chunks)} chunks from {file.filename}"}

@app.post("/chat")
async def chat_with_assistant(request: ChatRequest):
    chat_id = request.chat_id
    if not chat_id:
        new_chat = supabase.table("chats").insert({
            "user_email": request.user_email,
            "title": request.message[:40] + "..."
        }).execute()
        chat_id = new_chat.data[0]['id']

    supabase.table("messages").insert({
        "chat_id": chat_id,
        "role": "user",
        "content": request.message
    }).execute()

    result = gemini_client.models.embed_content(
        model="gemini-embedding-001",
        contents=request.message,
        config=types.EmbedContentConfig(output_dimensionality=768)
    )
    question_embedding = result.embeddings[0].values

    matching_docs = supabase.rpc("match_documents", {
        "query_embedding": question_embedding,
        "match_threshold": 0.70, 
        "match_count": 5,
        "p_user_email": request.user_email
    }).execute()

    context_text = "\n\n".join([f"Context: {doc['content']}" for doc in matching_docs.data])

    history = supabase.table("messages")\
        .select("*")\
        .eq("chat_id", chat_id)\
        .order("created_at")\
        .limit(10)\
        .execute()

    system_instruction = f"""
    You are Goofy AI, a severely caffeinated medical study assistant. 
    First, attempt to answer the user's question using the provided Context from their uploaded documents. 
    If the context DOES NOT contain the answer, use your general medical knowledge and the internet to answer, but briefly warn the user that you are pulling this from outside their uploaded notes.
    Maintain your witty, helpful persona.
    
    Context from uploaded documents:
    {context_text}
    """

    # 1. Attempt Primary Provider (Gemini)
    try:
        gemini_history = []
        for msg in history.data[:-1]: 
            gemini_history.append(
                types.Content(
                    role="model" if msg['role'] == "assistant" else "user",
                    parts=[types.Part.from_text(text=msg['content'])]
                )
            )
            
        chat_session = gemini_client.chats.create(
            model="gemini-3.6-flash", 
            config=types.GenerateContentConfig(
                system_instruction=system_instruction,
                tools=[{"google_search": {}}]
            ),
            history=gemini_history
        )
        response = chat_session.send_message(request.message)
        response_text = response.text

    # 2. Catch failures and fallback to Secondary Provider (OpenRouter)
    except Exception as e:
        logging.warning(f"Gemini API failed: {e}. Falling back to OpenRouter.")
        
        openrouter_messages = [{"role": "system", "content": system_instruction}]
        for msg in history.data[:-1]: 
            openrouter_messages.append({
                "role": msg['role'],
                "content": msg['content']
            })
        openrouter_messages.append({"role": "user", "content": request.message})
        
        # openrouter/free automatically routes to an available free model
        chat_completion = openrouter_client.chat.completions.create(
            messages=openrouter_messages,
            model="openrouter/free", 
        )
        response_text = chat_completion.choices[0].message.content

    supabase.table("messages").insert({
        "chat_id": chat_id,
        "role": "assistant",
        "content": response_text
    }).execute()

    return {
        "chat_id": chat_id,
        "response": response_text,
        "sources": [] 
    }

@app.post("/clear_data")
async def clear_user_data(request: DeleteRequest):
    supabase.table("documents").delete().eq("user_email", request.user_email).execute()
    user_chats = supabase.table("chats").select("id").eq("user_email", request.user_email).execute()
    chat_ids = [chat['id'] for chat in user_chats.data]
    if chat_ids:
        supabase.table("messages").delete().in_("chat_id", chat_ids).execute()
    supabase.table("chats").delete().eq("user_email", request.user_email).execute()
    return {"message": "Storage reclaimed successfully."}

@app.get("/user_chats")
async def get_user_chats(user_email: str):
    response = supabase.table("chats").select("id, title").eq("user_email", user_email).order("created_at", desc=True).execute()
    return {"chats": response.data}

@app.get("/chat_history/{chat_id}")
async def get_chat_history(chat_id: str):
    response = supabase.table("messages").select("role, content").eq("chat_id", chat_id).order("created_at").execute()
    return {"messages": response.data}

@app.delete("/chat/{chat_id}")
async def delete_chat(chat_id: str, user_email: str):
    supabase.table("messages").delete().eq("chat_id", chat_id).execute()
    supabase.table("chats").delete().eq("id", chat_id).eq("user_email", user_email).execute()
    return {"message": "Chat deleted"}

@app.get("/user_files")
async def get_user_files(user_email: str):
    response = supabase.table("documents").select("filename").eq("user_email", user_email).execute()
    filenames = list(set([doc['filename'] for doc in response.data]))
    return {"files": filenames}

@app.delete("/file")
async def delete_file(filename: str, user_email: str):
    supabase.table("documents").delete().eq("filename", filename).eq("user_email", user_email).execute()
    return {"message": "File deleted"}